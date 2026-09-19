"""Attachment readers — one interface for every format.

read_attachment(path) -> DocText(text, readable, meta)
`readable=False` feeds the `unreadable` escalation rather than raising.
"""
import json
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class DocText:
    text: str = ""
    readable: bool = True
    meta: dict = field(default_factory=dict)


def read_attachment(path: str | Path) -> DocText:
    p = Path(path)
    if not p.exists():
        return DocText(readable=False, meta={"error": "missing file"})
    if p.stat().st_size == 0:
        return DocText(readable=False, meta={"error": "empty file"})
    ext = p.suffix.lower()
    try:
        if ext == ".txt":
            return DocText(text=p.read_text(errors="replace"))
        if ext == ".pdf":
            return _read_pdf(p)
        if ext == ".docx":
            return _read_docx(p)
        if ext == ".xlsx":
            return _read_xlsx(p)
        return DocText(readable=False, meta={"error": f"unsupported ext {ext}"})
    except Exception as e:  # corrupt/garbled -> unreadable, not a crash
        return DocText(readable=False, meta={"error": f"{type(e).__name__}: {e}"})


def _read_pdf(p: Path) -> DocText:
    from pypdf import PdfReader

    r = PdfReader(str(p))
    text = "\n".join((page.extract_text() or "") for page in r.pages)
    if len(text.strip()) < 20:
        # image-only / scanned page — no text layer
        return DocText(text=text, readable=False, meta={"error": "no text layer (scanned)"})
    return DocText(text=text)


def _read_docx(p: Path) -> DocText:
    from docx import Document

    d = Document(str(p))
    lines = [para.text for para in d.paragraphs if para.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
            lines.append(": ".join(c for c in cells if c))
    return DocText(text="\n".join(lines))


def _read_xlsx(p: Path) -> DocText:
    import openpyxl

    wb = openpyxl.load_workbook(str(p), data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None and str(v).strip()]
            if vals:
                lines.append(": ".join(vals))
    return DocText(text="\n".join(lines))


def read_email_attachments(data_dir: str, email: dict) -> dict[str, DocText]:
    out = {}
    for rel in email.get("attachments") or []:
        out[rel] = read_attachment(Path(data_dir) / rel)
    return out


def load_email_record(data_dir: str, email_id: str) -> dict:
    p = Path(data_dir) / "inbox" / f"{email_id}.json"
    return json.loads(p.read_text())
